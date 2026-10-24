from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
D=Document(); sec=D.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.7);sec.page_height=Inches(8.3);sec.top_margin=Inches(.6);sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 23 if name=='Title' else 16 if name=='Heading 1' else 13)
 st.paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.12
D.core_properties.title='Bàn giao giao diện mobile DocPrep YHDP';D.core_properties.author='';D.core_properties.subject='Thiết kế chức năng điều kiện xử lý và hướng dẫn chạy FE'
def para(s,style=None): return D.add_paragraph(s,style)
def head(s): D.add_heading(s,level=1)
def table(headers,rows,widths):
 t=D.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'DEE4ED');borders.append(e)
 t._tbl.tblPr.append(borders)
 for i,h in enumerate(headers): t.rows[0].cells[i].text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cs=t.add_row().cells
  for i,txt in enumerate(row): cs[i].text=str(txt)
 for ri,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for k,v in [('top','75'),('bottom','75'),('left','100'),('right','100')]:
    x=OxmlElement('w:'+k);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');m.append(x)
   tcpr.append(m)
   if ri==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'244C87');tcpr.append(fill)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
    for r in p.runs:
     r.font.name='Calibri';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0);r.bold=ri==0
 D.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
D.add_paragraph('Bàn giao giao diện mobile DocPrep YHDP','Title')
para('Phiên bản triển khai ngày 16 tháng 10 năm 2026')
para('Bản FE mobile sử dụng bố cục đã chốt và các điều chỉnh mới về màu nhấn, khối học tập, avatar và linh vật AI. Các luồng học được nối với lớp API của web; chế độ trải nghiệm dùng dữ liệu lưu trên trình duyệt. Bảng dưới đây phân biệt rõ giao diện đã triển khai, xử lý có trong repo và phần máy chủ chưa hỗ trợ.')
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['home.jpg','profile.jpg']:
 f=ROOT/'docs/screenshots'/name
 if f.exists(): p.add_run().add_picture(str(f),width=Inches(1.7));p.add_run('       ')
para('Trang chủ với số liệu mẫu từ phiên học và Profile đổi avatar trực tiếp. Chuỗi ngày và tiến độ hiển thị theo dữ liệu thực của tài khoản trải nghiệm, không cố định 12 ngày hoặc 65%.')
D.add_page_break();head('1 Phạm vi và cách chạy')
para('FE nằm trong thư mục mobile và chạy bằng React cùng Vite. Đây là giao diện web tối ưu cho điện thoại; chưa đóng gói thành ứng dụng iOS hoặc Android native. Bản web yhdp frontend và snapshot repo gốc được giữ nguyên.')
table(['Nội dung','Bàn giao và điều kiện'],[
('Khởi chạy','Tại thư mục mobile chạy npm install rồi npm run dev. Mở http://127.0.0.1:5190. Chọn Trải nghiệm rồi chọn Người học, Cán bộ xem báo cáo hoặc Quản trị viên.'),
('Chế độ máy chủ','Màn Đăng nhập dùng mã kích hoạt cho người học và email mật khẩu cho cán bộ. FE gọi /api/v1 qua Vite proxy đến BACKEND_URL, mặc định http://127.0.0.1:8000. Máy chủ phải cung cấp /v1. Chưa kiểm thử tích hợp backend hoặc mô hình AI thật trong lần bàn giao này.'),
('Điều hướng','Năm mục cố định Trang chủ, Học tập, Thẻ ôn tập, Tiến độ, Profile. Thi thử, Luyện câu hỏi, Lịch ôn và Lộ trình nằm trong Học tập. Gia sư mở từ linh vật góc dưới.'),
('Phạm vi loại bỏ','Không mở các route sổ tay câu lưu, ca lâm sàng, thư viện tài liệu độc lập, trả lời ngắn, câu hỏi AI hạng B, sơ đồ khái niệm, workspace giảng viên hoặc thẩm định. Báo cáo ngân hàng tổng hợp vẫn giữ theo repo.'),
('Điều chỉnh mới','Profile được bổ sung theo quyết định thiết kế mới dù tài liệu phạm vi cũ đã bỏ trang tài khoản. Tên và cài đặt ôn tập dùng PATCH /me; giới tính minh họa và linh vật lưu theo tài khoản và chế độ trên thiết bị.'),
('Chú giải bảng đối chiếu','R nghĩa là repo có xử lý; M nghĩa là minh họa; N nghĩa là hợp đồng hoặc tính năng chưa có xử lý hoàn chỉnh. Các ký hiệu không có nghĩa đã nghiệm thu backend hoặc dữ liệu chuyên môn.'),
('Nguồn đối chiếu','Tài liệu 01 Phạm vi và chức năng trong yhdp frontend, bảng 142 mục RP đã trích xuất; web hiện có; repo Regnum771/NLHN-YHDP tại commit 5f32ff6a7d1c767a69f25cc7f215f67d502f232c; bản thiết kế V6 và các yêu cầu chỉnh tiếp theo trong trao đổi.')
],[1.65,8.75])
D.add_page_break();head('2 Thiết kế và thao tác trên mobile')
ui=[
('Trang chủ','Header xanh Navy, chữ DocPrep trắng, badge lửa nhỏ góc phải, thẻ chào đè nhẹ lên header.','GET /home và /session/today. Lời chào theo giờ; tên từ hồ sơ. Nam hoặc nữ theo lựa chọn, chưa chọn dùng hình trung tính. Chuỗi chỉ tăng khi có hoạt động học được ghi.'),
('Thẻ đến hạn','Icon Layers phẳng màu xanh, số thẻ đến hạn nổi bật, CTA Ôn thẻ ngay có gradient và bóng xanh nhẹ.','Số lượng đọc từ API. Khi không có thẻ thì dẫn đến bộ thẻ. Không dựng thẻ đến hạn giả hoặc dùng số 8 cố định.'),
('Phiên học hôm nay','Gom các bước thành một khối trắng với divider mờ. Icon thẻ màu Blue, điểm yếu Amber, học tiếp Teal.','Chỉ hiện các bước server trả; tối đa ba bước. Dẫn đúng bước qua stepPath, giữ trạng thái đã xong và tiếp tục bước chưa hoàn thành.'),
('Tiếp tục học','Danh sách bài gần đây với thanh tiến độ nhỏ dưới tên và dòng trạng thái.','Tiến độ lấy từ recent.progress. Giá trị null không đổi thành 0. Chạm bài mở đúng unit; số liệu thay đổi sau khi thực hiện bước học.'),
('Đăng nhập và quyền','Mã kích hoạt cho learner; email mật khẩu cho admin và ministry viewer. Trải nghiệm có ba vai trò tương ứng.','Không có OAuth, tự đăng ký hoặc quên mật khẩu người học. Token opaque và refresh do lớp session quản lý. Vai trò không hợp lệ bị chặn trước khi mở khu vực.'),
('Đồng thuận','Dịch vụ bắt buộc và nghiên cứu tùy chọn là hai checkbox riêng.','Version cũ yêu cầu đọc lại. Không chọn nghiên cứu vẫn tiếp tục học. Profile cho rút nghiên cứu hoặc dịch vụ; rút dịch vụ cần đồng ý lại để học.'),
('Chọn trọng tâm','Chọn nhóm học và vào bài đánh giá hoặc học ngay.','PATCH /me lưu focus module và cờ đã xử lý lời mời đánh giá. Chọn nghề sinh viên Y học dự phòng được nhận đúng phạm vi hỗ trợ; không cấp quyền mới.'),
('Học tập','Tìm không dấu, lọc có học liệu hoặc đã học; bộ lọc chi tiết mở khi cần.','Nhóm, vấn đề cốt lõi, nhóm nhận thức và trạng thái gửi theo hợp đồng /modules. Ba nhóm nhận thức gộp khác sáu bậc Bloom demo. Bộ lọc chưa có metadata trả rỗng, không suy diễn.'),
('Chi tiết học phần','Mở sheet danh sách unit theo thứ tự, chỉ rõ học liệu chưa có.','Link từ Tiến độ có hp mở đúng học phần. Không có học liệu thì hiện thông báo thay vì tạo nội dung.'),
('Bài học và quiz','Đọc nội dung, xem nguồn, chọn một phương án, bấm Kiểm tra rồi xem phản hồi.','Trước kiểm tra không tô đáp án chuẩn. Câu đã chấm không sửa chọn. Idempotency Key bảo toàn lần gửi khi retry. Nút và thứ tự bước do trạng thái API quyết định.'),
('Bài chữa và điều chỉnh','Hiển thị bài chữa hoặc ví dụ nếu response cung cấp, sau đó tiếp tục bài.','Không tự sinh bài chữa khi thiếu remediation. Tổng kết chỉ nói điều chỉnh mà server trả, không mặc định AI luôn thay độ khó.'),
('Luyện câu hỏi','Chọn chuyên khoa hoặc trộn ngẫu nhiên, chọn 5, 10 hoặc 20 câu, nhận giải thích sau từng câu.','Tạo bộ tùy chọn chạy demo. Backend gốc chưa có POST tạo bộ tùy ý nên nút tạo bị khóa ở live; vẫn mở bộ sẵn và phiên hôm nay. Không thay bằng câu nháp.'),
('Ôn thẻ','Hiện mặt trước, lật đáp án rồi đánh giá Học lại, Khó, Tốt hoặc Dễ.','POST review với Idempotency Key. Mốc xem trước từ thẻ API, không hardcode live. Hàng đợi được đọc lại sau đánh giá. Tạm ẩn thẻ qua PATCH.'),
('Bộ thẻ và duyệt thẻ','Tổng hợp theo học phần, tìm thẻ, mở nguồn, tạm ẩn hoặc mở lại thẻ được tải.','Live chỉ có endpoint thẻ đến hạn; chưa tải được toàn thư viện hoặc mọi thẻ đã ẩn. Demo cho duyệt kho mẫu. Không tự soạn, sửa nội dung hay xóa thẻ.'),
('Gia sư AI','Linh vật nhỏ bên phải phía trên footer, mặc định robot. Mở hội thoại hoặc đổi linh vật trong Profile.','Sáu lựa chọn từ page mascot, tải và phục vụ tại máy. Robot, mèo, cú, gấu trúc, nhà khoa học, cáo. Lựa chọn lưu tức thời theo tài khoản và chế độ.'),
('Hỏi đáp có nguồn','Hỏi tự do hoặc từ câu đã làm; hiển thị trạng thái xử lý, câu trả lời và đoạn nguồn.','Chỉ hiện các bước trace đã chạy. Phân biệt trả lời, cần làm rõ, thiếu bằng chứng, từ chối và lỗi hệ thống. Demo không gọi mô hình AI.'),
('Lịch sử và đề xuất AI','Mở lại hội thoại; tải thêm sessions và messages theo cursor nếu có.','Không gửi thêm khi lịch sử đang lỗi hoặc còn phần chưa tải. Giải thích thêm, hữu ích và đề xuất luyện hoặc thêm ngày mai dùng endpoint riêng; chưa bấm thì chưa áp dụng.'),
('Gợi ý khi luyện','Trước trả lời có gợi ý khi API cho phép; sau trả lời có giải thích và hỏi gia sư.','Guard của máy chủ vẫn quyết định. Trong bài thi không có nút hoặc mascot AI. Không dùng UI để khẳng định mọi cổng chống lộ đáp án đã được chứng minh an toàn.'),
('Thi thử','Xem đề, giới hạn lượt, quy định; bắt đầu hoặc tiếp tục attempt hiện có.','Lượt thi chụp câu và thứ tự. Bộ đếm dựa remaining của server; reload không tạo lại deadline. Lỗi mạng báo trạng thái lưu để thử lại.'),
('Làm bài thi','Một câu mỗi màn, đánh dấu, trước sau, danh sách câu và xác nhận nộp.','Tự lưu đáp án và cờ; chọn lại phương án đã chọn có thể bỏ chọn. Nộp đợi flush câu chưa lưu. Hết hạn theo server; nộp lặp không chấm thêm.'),
('Kết quả thi','Tổng đúng sai bỏ trống, thời gian, phần cần ôn, lọc câu và xem lại từng đáp án.','Chỉ dùng bài đã nộp. Không dự đoán đỗ. Câu bỏ trắng không được biến thành đáp án sai đã gửi để sinh thẻ giả.'),
('Tiến độ','Hai góc nhìn năng lực và nhịp học; độ phủ, ưu tiên, lịch sử điểm, trí nhớ và heatmap.','Phân biệt độ phủ, tỷ lệ đúng và recall. Null hiển thị thiếu dữ liệu. Ma trận đầy đủ có ở demo nhưng backend group grid còn null, có thông báo tương ứng.'),
('Lịch ôn và lộ trình','Xem thẻ dự kiến đến hạn, hoạt động theo ngày và đề xuất học.','Dự kiến không là cam kết lịch bất biến. Luồng lộ trình mở rộng ghi rõ phần chưa được máy chủ hỗ trợ; không tự cập nhật lịch nhiều ngày khi chưa chấp nhận.'),
('Thông báo','Lọc tất cả hoặc chưa đọc, đánh dấu đọc, bật nhắc trong ứng dụng, mở kết quả thi.','Suy ra từ thẻ đến hạn và hoạt động lưu trên thiết bị. Không phải push notification, email hoặc SMS do máy chủ gửi.'),
('Profile','Avatar bác sĩ đặt giữa dưới tiêu đề. Chọn Nam, Nữ hoặc Chưa chọn sẽ đổi ngay; bỏ banner xem trước lặp lại.','Lưu thay đổi lưu giới tính tại thiết bị; tên, ngày thi và giới hạn ôn dùng PATCH /me. Mascot lưu ngay. Avatar không suy đoán từ tên. Nút lưu dùng màu phẳng, ít nổi hơn CTA trang chủ.'),
('Quyền dữ liệu','Xuất dữ liệu cá nhân; rút đồng thuận; xác nhận yêu cầu xóa bằng cụm XÓA.','Không xóa chỉ vì đóng sheet. Chế độ demo chỉ thao tác dữ liệu mẫu. Chính sách lưu giữ và văn bản đồng thuận thật phải được hoàn thiện trước vận hành.'),
('Phản ánh nội dung','Báo lỗi từ câu luyện, câu xem lại, thẻ hoặc câu trả lời AI. Chọn lý do và mô tả.','Đúng schema /reports. Live 501 hiện chưa hỗ trợ và giữ nhập để thử lại. Demo lưu report draft riêng tại trình duyệt; không hiện thành hàng đợi cán bộ. Nội dung bài học chưa có content id trong view nên chưa có điểm gửi riêng.'),
('Báo cáo cán bộ','Header và bộ lọc tối ưu mobile; bảng rộng cuộn bên trong khối.','Cán bộ chỉ xem tổng hợp theo repo. Bốn bộ lọc kỳ, cơ sở, đối tượng, vùng. Chưa có quyền scope cơ sở được giao do FE tự quyết định.'),
('KPI và hoạt động','Giữ năm KPI, biểu đồ 12 tuần và quay lại sau 1, 4, 8 tuần.','Null và suppressed không đổi thành 0. Nhóm dưới ngưỡng bị ẩn. Demo chỉ có fixture mặc định, chọn phạm vi khác không dùng lại số quốc gia để giả dữ liệu lọc.'),
('Cơ sở và ngân hàng','Bảng cơ sở và ngân hàng toàn hệ thống giữ công thức và nhãn gốc.','Ngân hàng không theo bộ lọc trường đối tượng vùng. Giữ trạng thái duyệt, độ phủ 54 vấn đề, trung vị vòng 1 vòng 2 và tỷ lệ cùng kết luận. Không mở workspace thẩm định.'),
('An toàn và tuân thủ','Hiện phiên bản kho và mô hình nếu có, tỷ lệ khẳng định được giữ, từ chối, gắn cờ, báo lỗi và P95.','Tỷ lệ trích dẫn dùng claims kept chia claims checked theo repo. Không suy ra đó là độ đúng y khoa; không thêm uptime hoặc trạng thái sao lưu giả.'),
('Cảnh báo và chi phí','Cảnh báo kích hoạt thấp, vòng 2 chậm và báo lỗi tồn. Chi phí có nhãn minh họa rõ.','Giữ ngưỡng repo. Chi phí client là tiền mẫu, không phải chi phí vận hành được đo. Không xuất khối minh họa trong Excel của server.'),
('Xuất báo cáo','Nút xuất dùng cùng bộ lọc đang chọn.','Live gọi /admin/dashboard/export với format xlsx; lấy blob từ máy chủ và không xuất định danh learner. Demo xuất JSON có nhãn minh họa. Chưa kiểm thử tệp live khi chưa có backend.'),
('Nhập tài khoản','Admin chọn hoặc nhập CSV, xem lỗi từng dòng, chỉ xác nhận khi hợp lệ.','Demo tạo DEMO LOCAL, không dùng để đăng nhập thật hoặc gửi thư. Import preview và confirm của backend vẫn là 501; UI không khẳng định đã cấp tài khoản thật.'),
('Hạ tầng và xử lý nền','FE giữ hợp đồng API, session refresh, lỗi truy cập và các trạng thái chờ.','Elo, FSRS, truy hồi, hậu kiểm, migration, audit và công cụ seed thuộc backend. Không biến các dòng kỹ thuật trong Word thành nút quản trị giả.')
]
table(['Màn hình hoặc chức năng','Bố cục và thao tác','Dữ liệu và điều kiện'],ui,[1.6,3.4,5.4])
head('3 Đối chiếu đầy đủ 142 mục chức năng')
para('Bảng giữ mô tả và điều kiện của từng mã RP trong tài liệu nguồn để đội triển khai kiểm tra nghiệp vụ. Dòng Không gian mobile chỉ điểm vào hoặc nơi xử lý. Những mục thuộc backend không được coi là tính năng mới đã hoàn thành chỉ vì có giao diện. Riêng giới hạn Profile trong RP F01 được thay bằng yêu cầu Profile mới đã chốt ở mục 2.')
routes={'A':'Đăng nhập và Profile','F':'Chọn trọng tâm và Profile','H':'Học tập','P':'Phản hồi và tổng kết','D':'Trang chủ và Phiên học','L':'Luyện câu hỏi','C':'Thẻ ôn tập','Q':'Gia sư AI','T':'Gợi ý và đề xuất','G':'Trace AI và backend','E':'Thi thử và kết quả','S':'Tiến độ và Lịch ôn','B':'Báo cáo cán bộ','V':'Phản ánh hoặc Cấp mã','O':'Backend và trạng thái API'}
req=json.loads((ROOT/'docs/source-requirements.json').read_text())
rows=[]
for ident,title,flow,limits in req:
 group=ident.split('-')[1][0]
 mobile=routes[group]
 if ident in ['RP-V01','RP-V04','RP-V09']:mobile='Backend nội dung'
 if ident=='RP-A14':mobile='Công cụ kỹ thuật backend'
 if ident=='RP-V07':mobile='Báo lỗi nội dung'
 if ident=='RP-V08':mobile='Quản trị vào Cấp mã'
 rows.append((ident+' '+title+'\nMobile '+mobile,flow,limits))
table(['Mã chức năng và không gian mobile','Luồng xử lý theo repo và Word','Điều kiện và giới hạn theo nguồn'],rows,[2.35,4.0,4.05])
head('4 Kiểm tra và giới hạn bàn giao')
para('FE đã được build production, chạy lint và thực hiện 77 kiểm tra logic bằng bộ test của dự án cùng các kiểm tra bổ sung cho mặc định avatar, tách dữ liệu tài khoản và lọc học phần. Kiểm tra trình duyệt dùng Codex In app Browser với các viewport điện thoại khoảng 355 đến 429 CSS px và khung desktop. Bảng dưới ghi kết quả thao tác đã quan sát, không thay thế nghiệm thu backend.')
table(['Luồng','Kết quả quan sát'],[
('Người học','Đăng nhập demo, đồng ý dịch vụ mà không đồng ý nghiên cứu, chọn trọng tâm, mở bài học. Quiz chọn và chấm đúng sai, mở nguồn, tiếp tục bước.'),
('Ôn thẻ','Lật đáp án và gửi mức Tốt, hàng đợi giảm một thẻ, badge chuỗi từ 0 thành 1 ngày sau hoạt động học.'),
('Profile và mascot','Nam đổi avatar trực tiếp; Nữ cùng tên mẫu được lưu và hiện ở trang chủ. Chọn mèo vẫn giữ sau reload; đã đưa lại Robot trước bàn giao.'),
('Gia sư','Câu hỏi mẫu tạo hội thoại, phản hồi có nguồn và nhãn demo không gọi AI. Ô nhập nằm phía trên footer. Lịch sử có nút tải tiếp khi API có cursor.'),
('Luyện và phản ánh','Tạo bộ 5 câu, chọn sai thấy đáp án chuẩn và giải thích; thẻ sinh từ câu sai. Lưu phản ánh mẫu có thông báo chỉ lưu tại trình duyệt.'),
('Thi thử','Tạo bài đánh giá, trả lời và đánh dấu, reload giữ trạng thái; timer tiếp tục. Mở danh sách, nộp khi còn câu trắng, xem kết quả. Không có nav học hoặc AI trong màn làm bài.'),
('Cán bộ và admin','Xem đủ các khối báo cáo, chọn nhóm nhỏ thấy số liệu bị ẩn. Xuất được JSON mẫu. CSV sai audience bị chặn; CSV hợp lệ tạo mã mẫu sau xác nhận.'),
('Chưa nghiệm thu','API live, token đăng nhập thật, xuất Excel thật, lưu hồ sơ trên backend, gửi báo lỗi thật, import tài khoản thật, mô hình AI và DB chưa được chạy tích hợp. Không triển khai hoặc xuất bản lên Internet.')
],[2.05,8.35])
para('Tài liệu kỹ thuật và bằng chứng kiểm tra: README.md, docs/QA.md, docs/source-requirements.json, docs/screenshots và public/mobile/sources.json. Bản thiết kế gốc để đối chiếu nằm trong docs/design/approved-v6. Sprite linh vật sử dụng giấy phép MIT của nilbuild/page-mascot, bản giấy phép được giữ trong public/mobile.')
# Black headings, no inherited title borders, clean page-number footer.
for e in list(D.element.xpath('//w:pBdr')): e.getparent().remove(e)
for st in D.styles:
 for e in list(st.element.xpath('.//w:pBdr')): e.getparent().remove(e)
for style in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 D.styles[style].font.color.rgb=RGBColor(0,0,0)
f=D.sections[0].footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('DocPrep YHDP   ');r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
out=ROOT/'docs/DocPrep_YHDP_Ban_giao_FE_mobile.docx';D.save(out);print(out)
