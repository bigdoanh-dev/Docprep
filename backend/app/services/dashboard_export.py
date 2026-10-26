"""The ministry overview as an .xlsx file, built from the same view the page shows (FR-ADM-16).

One sheet per block, in the page's order. Every sheet opens with what the figures are: the period,
the filters, when they were computed, the note on small groups, and "Dữ liệu mẫu" when the view holds
synthetic learners. Figures are numbers; shares are fractions with a percent format; a withheld
figure is an empty cell.
"""
from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from app.config import settings

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TITLE = "Tổng quan toàn quốc"
SUPPRESSED = "Không đủ học viên để hiển thị số liệu cho bộ lọc này."
PERCENT = "0.0%"
DAYS = "0.0"
AUDIENCES = {"student": "Sinh viên", "graduate": "Đã tốt nghiệp"}


def file_name(view: dict) -> str:
    return f"tong-quan-{view['period']}.xlsx"


def _computed_at(view: dict) -> str:
    at = view["computed_at"]
    at = datetime.fromisoformat(at.replace("Z", "+00:00")) if isinstance(at, str) else at
    return at.astimezone(ZoneInfo(settings()["local_timezone"])).strftime("%d/%m/%Y %H:%M")


def _school_label(view: dict) -> str:
    code = view["selected"]["school"]
    if code is None:
        return "Tất cả"
    return next((o["label_vi"] for o in view["filters"]["schools"] if o["value"] == code), code)


def _heading(view: dict) -> list[list]:
    """The lines above the figures of a sheet."""
    selected = view["selected"]
    year, month = view["period"].split("-")
    rows = [[TITLE], ["Kỳ", f"Tháng {int(month)}/{year}"],
            ["Cơ sở đào tạo", _school_label(view)],
            ["Đối tượng", AUDIENCES.get(selected["audience"], "Tất cả")],
            ["Vùng", selected["region"] or "Toàn quốc"], [f"Tính lúc {_computed_at(view)}"],
            [f"Nhóm dưới {view['suppression_min']} học viên được ẩn"]]
    if view["sample_data"]:
        rows.append(["Dữ liệu mẫu"])
    return rows + [[]]


class _Sheet:
    """A sheet being written: rows of cells, each cell optionally with a number format."""

    def __init__(self, book: Workbook, title: str, view: dict, first: bool = False):
        self.sheet = book.active if first else book.create_sheet()
        self.sheet.title = title
        for row in _heading(view):
            self.sheet.append(row)

    def row(self, *cells, formats: dict[int, str] | None = None) -> None:
        self.sheet.append(list(cells))
        for column, number_format in (formats or {}).items():
            cell = self.sheet.cell(row=self.sheet.max_row, column=column + 1)
            if cell.value is not None:
                cell.number_format = number_format

    def figure(self, label: str, value, number_format: str | None = None) -> None:
        self.row(label, value, formats={1: number_format} if number_format else None)


def _overview(sheet: _Sheet, kpis: dict) -> None:
    median = kpis["mock_median"]
    sheet.figure("Học viên hoạt động trong tháng", kpis["active_learners"])
    sheet.figure("Học viên hoạt động tháng trước", kpis["active_learners_previous"])
    sheet.figure("Tài khoản đã cấp", kpis["accounts_issued"])
    sheet.figure("Tài khoản đã kích hoạt", kpis["accounts_activated"])
    sheet.figure("Câu luyện tập đã làm", kpis["practice_items"])
    sheet.figure("Lượt thi thử", kpis["mock_attempts"])
    sheet.figure("Trung vị điểm thi thử (tỷ lệ câu đúng)", median and median["share"], PERCENT)
    sheet.figure("Số câu của đề", median and median["form_size"])
    sheet.figure("Số học viên (có đồng ý nghiên cứu)", median and median["learners"])
    sheet.figure("Câu trả lời AI có trích dẫn", kpis["cited_claims_share"], PERCENT)
    sheet.figure("Mục tiêu trích dẫn", kpis["cited_claims_target"], PERCENT)


def _activity(sheet: _Sheet, weekly: dict) -> None:
    sheet.row("Tuần bắt đầu", "Học viên hoạt động")
    for week in weekly["weeks"]:
        start = week["week_start"]
        sheet.row(date.fromisoformat(start) if isinstance(start, str) else start, week["active_learners"],
                  formats={0: "dd/mm/yyyy"})
    sheet.row()
    sheet.row("Quay lại sau", "Tỷ lệ", "Số học viên")
    for item in weekly["retention"]:
        sheet.row(f"{item['weeks']} tuần", item["share"], item["learners"], formats={1: PERCENT})


def _alerts(sheet: _Sheet, alerts: list[dict]) -> None:
    if not alerts:
        sheet.row("Không có mục nào cần chú ý")
    for alert in alerts:
        sheet.row(alert["title_vi"], alert["detail_vi"])


def _accuracy(sheet: _Sheet, accuracy: dict, low_share: float) -> None:
    sheet.figure("Học viên có câu trả lời thi thử (có đồng ý nghiên cứu)", accuracy["learners"])
    sheet.row()
    sheet.row("Học phần", "Tỷ lệ đúng", f"Thấp (dưới {low_share:.0%})", "Số học viên")
    for item in accuracy["items"]:
        sheet.row(item["short_name"], item["share"], "Có" if item["low"] else "", item["learners"],
                  formats={1: PERCENT})


def _bank(sheet: _Sheet, bank: dict) -> None:
    sheet.row("Toàn bộ ngân hàng: không theo bộ lọc cơ sở, đối tượng, vùng")
    sheet.figure("Đã duyệt", bank["approved"])
    sheet.figure("Đang duyệt", bank["in_review"])
    sheet.figure("Nháp", bank["draft"])
    sheet.figure("Loại", bank["rejected"])
    sheet.figure(f"Vấn đề cốt lõi có ≥ {bank['min_approved_per_issue']} câu đã duyệt", bank["issues_covered"])
    sheet.figure("Tổng số vấn đề cốt lõi", bank["issues_total"])
    sheet.figure("Duyệt vòng 1, trung vị (ngày)", bank["round1_median_days"], DAYS)
    sheet.figure("Duyệt vòng 2, trung vị (ngày)", bank["round2_median_days"], DAYS)
    sheet.figure("Mục tiêu vòng 2 (ngày)", bank["round2_target_days"])
    sheet.figure("Hai vòng cùng kết luận", bank["rounds_agree_share"], PERCENT)
    sheet.figure("Báo lỗi từ học viên nhận được", bank["reports_received"])
    sheet.figure("Báo lỗi đã xử lý", bank["reports_resolved"])


def _ai(sheet: _Sheet, ai: dict) -> None:
    sheet.figure("Câu trả lời", ai["answers"])
    sheet.figure("Phiên bản mô hình", ", ".join(ai["model_versions"]))
    sheet.figure("Phiên bản kho tài liệu", ", ".join(ai["corpus_versions"]))
    sheet.figure("Khẳng định có trích dẫn đến trang", ai["cited_claims_share"], PERCENT)
    sheet.figure("Mục tiêu", ai["cited_claims_target"], PERCENT)
    sheet.figure("Từ chối vì thiếu bằng chứng", ai["insufficient_share"], PERCENT)
    sheet.figure("Từ chối tư vấn ca bệnh cụ thể", ai["declined_share"], PERCENT)
    sheet.figure("Bị hậu kiểm trích dẫn gắn cờ", ai["flagged_share"], PERCENT)
    sheet.figure("Báo lỗi trên 1.000 câu trả lời", ai["reports_per_1000"], DAYS)


def _schools(sheet: _Sheet, schools: dict) -> None:
    sheet.row("Cơ sở", "Vùng", "Tài khoản đã cấp", "Đã kích hoạt", "Tỷ lệ kích hoạt", "Kích hoạt thấp",
              "Hoạt động mỗi tuần (trung bình)", "Lượt thi thử", "Trung vị điểm (tỷ lệ đúng)", "Số câu của đề",
              "Số học viên", "Thay đổi từ lần đầu đến gần nhất", "Số học viên")
    for row in schools["rows"]:
        median, change = row["mock_median"], row["first_to_latest"]
        issued = row["accounts_issued"]
        sheet.row(row["name"], row["region"], issued, row["accounts_activated"],
                  row["accounts_activated"] / issued if issued else None, "Có" if row["activation_low"] else "",
                  row["weekly_active"], row["mock_attempts"], median and median["share"],
                  median and median["form_size"], median and median["learners"], change and change["share"],
                  change and change["learners"], formats={4: PERCENT, 6: DAYS, 8: PERCENT, 11: PERCENT})
    if schools["hidden"]:
        sheet.row()
        sheet.row("Cơ sở dưới ngưỡng (ẩn số liệu)", "Vùng", "Số học viên")
        for school in schools["hidden"]:
            sheet.row(school["name"], school["region"], school["learners"])


def _compliance(sheet: _Sheet, compliance: dict) -> None:
    sheet.figure("Độ trễ một lượt hỏi đáp, P95 (ms)", compliance["chat_p95_ms"])
    sheet.figure("Mục tiêu (ms)", compliance["chat_target_ms"])
    sheet.figure("Học viên đồng ý dùng dữ liệu cho nghiên cứu", compliance["research_consent_learners"])
    sheet.figure("Tỷ lệ trên học viên đã kích hoạt", compliance["research_consent_share"], PERCENT)
    sheet.figure("Rút đồng ý", compliance["consent_withdrawals"])
    sheet.figure("Yêu cầu xóa dữ liệu", compliance["deletion_requests"])
    sheet.figure("Thời gian xóa, trung vị (ngày)", compliance["deletion_median_days"], DAYS)


def build(view: dict) -> bytes:
    """The workbook for a view (an `AdminOverview` as a dict)."""
    book = Workbook()
    first = _Sheet(book, "Tổng quan", view, first=True)
    if view["suppressed"] or view["kpis"] is None:
        first.row(SUPPRESSED)
    else:
        low_share = settings()["dashboard"]["low_share"]
        _overview(first, view["kpis"])
        _activity(_Sheet(book, "Hoạt động", view), view["weekly_active"])
        _alerts(_Sheet(book, "Cần chú ý", view), view["alerts"])
        _accuracy(_Sheet(book, "Câu đúng theo miền", view), view["accuracy_by_module"], low_share)
        _bank(_Sheet(book, "Ngân hàng câu hỏi", view), view["question_bank"])
        _ai(_Sheet(book, "An toàn AI", view), view["ai_safety"])
        _schools(_Sheet(book, "Theo cơ sở", view), view["schools"])
        _compliance(_Sheet(book, "Vận hành", view), view["compliance"])
    out = BytesIO()
    book.save(out)
    return out.getvalue()
