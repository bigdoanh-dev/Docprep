"""The overview as an Excel file: the same view as the screen, one sheet per block (FR-ADM-16)."""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import _test_conn

EXPORT = "/v1/admin/dashboard/export"
SHEETS = ["Tổng quan", "Hoạt động", "Cần chú ý", "Câu đúng theo miền", "Ngân hàng câu hỏi", "An toàn AI",
          "Theo cơ sở", "Vận hành"]


def _download(client, role="ministry_viewer", **filters):
    return client.get(EXPORT, params={"period": "2026-09", **filters}, headers=client.role_headers[role])


def _first_cells(sheet) -> list:
    return [row[0] for row in sheet.iter_rows(values_only=True)]


def _value(sheet, label: str):
    return next(row[1] for row in sheet.iter_rows(values_only=True) if row[0] == label)


def test_the_file_has_a_sheet_per_block_each_saying_what_it_shows(cohort_api):
    r = _download(cohort_api)
    assert r.status_code == 200, r.text
    assert r.headers["content-disposition"] == 'attachment; filename="tong-quan-2026-09.xlsx"'
    book = load_workbook(BytesIO(r.content))
    assert book.sheetnames == SHEETS
    for sheet in book.worksheets:
        assert sheet["A1"].value == "Tổng quan toàn quốc", sheet.title
        assert "Dữ liệu mẫu" in _first_cells(sheet), sheet.title
        assert "Nhóm dưới 10 học viên được ẩn" in _first_cells(sheet), sheet.title
        assert any(str(cell).startswith("Tính lúc 13/10/2026 10:00") for cell in _first_cells(sheet)), sheet.title


def test_the_file_and_the_screen_agree(cohort_api):
    view = cohort_api.get("/v1/admin/dashboard/overview", params={"period": "2026-09"},
                          headers=cohort_api.role_headers["ministry_viewer"]).json()
    book = load_workbook(BytesIO(_download(cohort_api).content))
    assert _value(book["Tổng quan"], "Học viên hoạt động trong tháng") == view["kpis"]["active_learners"]
    assert _value(book["Ngân hàng câu hỏi"], "Đã duyệt") == view["question_bank"]["approved"]
    hidden = view["schools"]["hidden"][0]
    assert any(row[:3] == (hidden["name"], hidden["region"], hidden["learners"])
               for row in book["Theo cơ sở"].iter_rows(values_only=True))


def test_a_share_is_a_number_with_a_percent_format(cohort_api):
    sheet = load_workbook(BytesIO(_download(cohort_api).content))["An toàn AI"]
    cell = next(row[1] for row in sheet.iter_rows() if row[0].value == "Khẳng định có trích dẫn đến trang")
    assert isinstance(cell.value, (int, float)) and 0 <= cell.value <= 1     # a share of 1.0 reads back as 1
    assert "%" in cell.number_format


def test_a_suppressed_view_is_one_sheet_with_the_notice(cohort_api):
    book = load_workbook(BytesIO(_download(cohort_api, school="CSDT-F").content))
    assert book.sheetnames == ["Tổng quan"]
    assert "Không đủ học viên để hiển thị số liệu cho bộ lọc này." in _first_cells(book.active)


def test_each_download_is_written_to_the_audit_log(cohort_api):
    def rows():
        with _test_conn() as c:
            return c.execute("select actor_type, actor_id, target, details from audit_log "
                             "where action = 'dashboard_export' order by id").fetchall()

    before = len(rows())
    assert _download(cohort_api, "admin", region="Miền Nam").status_code == 200
    after = rows()
    assert len(after) == before + 1
    last = after[-1]
    assert (last["actor_type"], last["actor_id"]) == ("staff", str(cohort_api.staff_ids["admin"]))
    assert last["details"] == {"period": "2026-09", "school": None, "audience": None, "region": "Miền Nam"}


def test_a_learner_cannot_download_it(cohort_api):
    assert _download(cohort_api, "learner").status_code == 403
